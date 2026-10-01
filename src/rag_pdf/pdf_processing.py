"""PDF validation, text extraction, and conservative cleanup."""

from __future__ import annotations

import hashlib
import math
import re
from collections import Counter
from pathlib import Path

import pymupdf
from langdetect import DetectorFactory, detect_langs
from langdetect.lang_detect_exception import LangDetectException

from rag_pdf.config import PDFRules
from rag_pdf.errors import ValidationError
from rag_pdf.models import ExtractedDocument, PageText

DetectorFactory.seed = 0


def extract_document(
    data: bytes,
    filename: str,
    rules: PDFRules | None = None,
) -> ExtractedDocument:
    """Extract a supported document while preserving its useful source location."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(data, filename, rules)
    if suffix == ".docx":
        return extract_docx(data, filename, rules)
    raise ValidationError("Only PDF and Word (.docx) files are accepted.")


def extract_docx(data: bytes, filename: str, rules: PDFRules | None = None) -> ExtractedDocument:
    """Extract Word paragraphs and tables as labeled, searchable text."""
    from io import BytesIO

    from docx import Document

    rules = rules or PDFRules()
    if not data:
        raise ValidationError("The uploaded Word document is empty.")
    if len(data) > rules.max_file_bytes:
        raise ValidationError("The Word document is larger than the 20 MB limit.")
    if Path(filename).suffix.lower() != ".docx":
        raise ValidationError("Only modern Word (.docx) files are accepted.")
    try:
        word = Document(BytesIO(data))
    except Exception as exc:
        raise ValidationError("The uploaded file could not be opened as a Word document.") from exc

    # Walk the document body in order, retaining tables as row/column labels so values
    # stay associated with their headings during retrieval and answer generation.
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    blocks: list[str] = []
    for child in word.element.body.iterchildren():
        if isinstance(child, CT_P):
            text = Paragraph(child, word).text.strip()
            if text:
                blocks.append(text)
        elif isinstance(child, CT_Tbl):
            table = Table(child, word)
            rows = [[" ".join(cell.text.split()) for cell in row.cells] for row in table.rows]
            rows = [row for row in rows if any(row)]
            if rows:
                headers = rows[0]
                blocks.append("Table:")
                if len(rows) > 1 and len(headers) > 1:
                    for row in rows[1:]:
                        entries = [
                            f"{headers[i]}: {value}"
                            for i, value in enumerate(row)
                            if value and i < len(headers)
                        ]
                        blocks.append("; ".join(entries))
                else:
                    blocks.extend(" | ".join(row) for row in rows)

    text = "\n".join(blocks)
    character_count = len(re.sub(r"\s", "", text))
    if character_count < max(100, rules.min_total_characters // 10):
        raise ValidationError("Very little text could be extracted from the Word document.")
    language, confidence = _detect_language([text])
    if language != "en" or confidence < rules.min_english_probability:
        raise ValidationError(
            "The extracted document is not confidently English "
            f"(detected {language!r}, confidence {confidence:.0%})."
        )
    return ExtractedDocument(
        document_id=document_hash(data), filename=Path(filename).name,
        page_count=1, pages=(PageText(page_number=1, text=text),),
        language=language, language_confidence=confidence, character_count=character_count,
    )


def document_hash(data: bytes) -> str:
    """Return a stable ID used to isolate indexes and chat histories."""

    return hashlib.sha256(data).hexdigest()


def extract_pdf(
    data: bytes,
    filename: str,
    rules: PDFRules | None = None,
) -> ExtractedDocument:
    """Validate and extract a text-only English PDF page by page."""

    rules = rules or PDFRules()
    _validate_file(data, filename, rules)

    try:
        document = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise ValidationError("The uploaded file could not be opened as a PDF.") from exc

    try:
        if document.needs_pass:
            raise ValidationError("Password-protected PDFs are not supported.")

        page_count = document.page_count
        if not rules.min_pages <= page_count <= rules.max_pages:
            raise ValidationError(
                f"The PDF has {page_count} pages; it must contain "
                f"{rules.min_pages}–{rules.max_pages} pages."
            )

        raw_pages = [page.get_text("text", sort=True) for page in document]
    finally:
        document.close()

    cleaned_pages = _clean_pages(raw_pages)
    character_count = sum(len(re.sub(r"\s", "", text)) for text in cleaned_pages)
    text_page_count = sum(
        len(re.sub(r"\s", "", text)) >= rules.min_characters_per_text_page
        for text in cleaned_pages
    )

    if character_count < rules.min_total_characters:
        raise ValidationError(
            "Very little text could be extracted. Please use a text-based PDF, not a scan."
        )

    required_text_pages = math.ceil(page_count * rules.min_text_page_ratio)
    if text_page_count < required_text_pages:
        raise ValidationError(
            "Most pages do not contain extractable text. Scanned PDFs are not supported."
        )

    language, confidence = _detect_language(cleaned_pages)
    if language != "en" or confidence < rules.min_english_probability:
        raise ValidationError(
            "The extracted document is not confidently English "
            f"(detected {language!r}, confidence {confidence:.0%})."
        )

    pages = tuple(
        PageText(page_number=index + 1, text=text)
        for index, text in enumerate(cleaned_pages)
        if text.strip()
    )
    return ExtractedDocument(
        document_id=document_hash(data),
        filename=Path(filename).name,
        page_count=page_count,
        pages=pages,
        language=language,
        language_confidence=confidence,
        character_count=character_count,
    )


def _validate_file(data: bytes, filename: str, rules: PDFRules) -> None:
    if Path(filename).suffix.lower() != ".pdf":
        raise ValidationError("Only PDF files are accepted.")
    if not data:
        raise ValidationError("The uploaded PDF is empty.")
    if len(data) > rules.max_file_bytes:
        max_mb = rules.max_file_bytes // (1024 * 1024)
        raise ValidationError(f"The PDF is larger than the {max_mb} MB limit.")
    if b"%PDF-" not in data[:1024]:
        raise ValidationError("The file extension is PDF, but its contents are not a valid PDF.")


def _clean_pages(raw_pages: list[str]) -> list[str]:
    normalized = [_normalize_lines(text) for text in raw_pages]
    repeated_margins = _find_repeated_margins(normalized)
    return [_remove_repeated_margins(text, repeated_margins) for text in normalized]


def _normalize_lines(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\u00ad", "")
    text = re.sub(r"(?<=\w)-\n(?=[a-z])", "", text)
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def _margin_key(line: str) -> str:
    """Normalize page numbers so 'Page 1 of 10' matches 'Page 2 of 10'."""

    return re.sub(r"\d+", "#", line.casefold()).strip()


def _find_repeated_margins(pages: list[str]) -> set[str]:
    page_candidates: list[set[str]] = []
    for page in pages:
        lines = [line for line in page.splitlines() if line]
        margin_lines = lines[:2] + lines[-2:] if lines else []
        page_candidates.append({_margin_key(line) for line in margin_lines if len(line) <= 160})

    counts = Counter(key for candidates in page_candidates for key in candidates)
    threshold = max(3, math.ceil(len(pages) * 0.60))
    return {key for key, count in counts.items() if key and count >= threshold}


def _remove_repeated_margins(text: str, repeated: set[str]) -> str:
    if not repeated:
        return text
    lines = text.splitlines()
    nonempty_indexes = [index for index, line in enumerate(lines) if line]
    margin_indexes = set(nonempty_indexes[:2] + nonempty_indexes[-2:])
    kept = [
        line
        for index, line in enumerate(lines)
        if not (index in margin_indexes and _margin_key(line) in repeated)
    ]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(kept)).strip()


def _detect_language(pages: list[str]) -> tuple[str, float]:
    sample = " ".join(page for page in pages if page)[:20_000]
    try:
        candidates = detect_langs(sample)
    except LangDetectException as exc:
        raise ValidationError("The document language could not be determined.") from exc

    if not candidates:
        raise ValidationError("The document language could not be determined.")
    best = candidates[0]
    return best.lang, float(best.prob)
