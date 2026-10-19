"""Native PDF and DOCX extraction, retaining tables and honest source locations."""

from __future__ import annotations

import io
import math
import re
import unicodedata
import zipfile
from pathlib import Path

from rag_pdf.errors import ValidationError
from rag_pdf.models import DocumentTable, ExtractedDocument, PageText
from rag_pdf.pdf_processing import document_hash

MAX_BYTES = 20 * 1024 * 1024
MAX_WORKSHEET_CELLS = 500_000


def make_table(table_id, data, location, page_number=0, caption="", headers=None):
    rows = [
        [unicodedata.normalize("NFKC", str(cell or "")).replace("\n", " ").strip() for cell in row]
        for row in data
    ]
    rows = [row for row in rows if any(row)]
    if not rows:
        return None
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    if headers is None:
        headers, rows = rows[0], rows[1:]
    headers = [str(h or f"Column {i + 1}") for i, h in enumerate(headers)]
    if not rows:
        return None
    return DocumentTable(
        table_id,
        tuple(unicodedata.normalize("NFKC", header) for header in headers),
        tuple(map(tuple, rows)),
        location,
        page_number,
        unicodedata.normalize("NFKC", caption),
    )


def extract_document(data: bytes, filename: str) -> ExtractedDocument:
    suffix = Path(filename).suffix.lower()
    extractors = {
        ".pdf": _pdf,
        ".docx": _docx,
        ".xlsx": _xlsx,
        ".pptx": _pptx,
    }
    if suffix not in extractors:
        raise ValidationError("Upload a PDF, Word .docx, Excel .xlsx, or PowerPoint .pptx file.")
    if not data or len(data) > MAX_BYTES:
        raise ValidationError("Documents must be nonempty and at most 20 MB.")
    try:
        pages, tables, warnings, page_count = extractors[suffix](data)
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError(f"Could not read {Path(filename).name}: {exc}") from exc
    count = sum(len(p.text) for p in pages) + sum(len(t.render()) for t in tables)
    if not count:
        raise ValidationError(
            "No native text or tables found. Image-only scans and image-based slide content "
            "are not supported."
        )
    # Numeric tables and short documents must not fail an English-language heuristic.
    return ExtractedDocument(
        document_hash(data),
        Path(filename).name,
        page_count,
        tuple(pages),
        "und",
        0.0,
        count,
        tuple(tables),
        tuple(warnings),
        suffix[1:],
    )


def _check_office_archive(data):
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 50_000 or sum(item.file_size for item in entries) > 100 * 1024 * 1024:
                raise ValidationError("The expanded Office document exceeds its size limit.")
    except zipfile.BadZipFile as exc:
        raise ValidationError("This is not a valid Office document.") from exc


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, (float,)) and not math.isfinite(value):
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value).replace("\n", " ").strip()


def _pdf(data):
    import pymupdf

    if b"%PDF-" not in data[:1024]:
        raise ValidationError("This file is not a PDF.")
    pages, tables, warnings = [], [], []
    with pymupdf.open(stream=data, filetype="pdf") as document:
        if document.needs_pass:
            raise ValidationError("Password-protected PDFs are not supported.")
        if document.page_count > 200:
            raise ValidationError("PDFs are limited to 200 pages.")
        for number, page in enumerate(document, 1):
            bounds = []
            try:
                found = page.find_tables()
                for table in found.tables:
                    table_id = f"p{number}-t{len(tables) + 1}"
                    surrounding = page.get_text(
                        clip=pymupdf.Rect(
                            0, max(0, table.bbox[1] - 45), page.rect.width, table.bbox[1]
                        )
                    ).strip()
                    external = table.header.external
                    parsed = make_table(
                        table_id,
                        table.extract(),
                        f"p. {number}, Table {len(tables) + 1}",
                        number,
                        surrounding,
                        headers=table.header.names if external else None,
                    )
                    if parsed:
                        tables.append(parsed)
                        bounds.append(pymupdf.Rect(table.bbox))
            except Exception as exc:
                warnings.append(
                    f"Page {number}: table detection failed ({type(exc).__name__}); "
                    "check the text preview."
                )
            # Exclude detected cells from prose so table values are not indexed twice.
            lines = []
            for block in page.get_text("dict", sort=True)["blocks"]:
                for line in block.get("lines", []):
                    text = "".join(
                        span["text"]
                        for span in line["spans"]
                        if not any(
                            box.contains(
                                pymupdf.Rect(span["bbox"]).tl
                                + (pymupdf.Rect(span["bbox"]).br - pymupdf.Rect(span["bbox"]).tl)
                                / 2
                            )
                            for box in bounds
                        )
                    ).strip()
                    if text:
                        lines.append(text)
            text = unicodedata.normalize("NFKC", "\n".join(lines))
            if text:
                pages.append(PageText(number, text, f"p. {number}"))
            elif not bounds:
                warnings.append(f"Page {number}: no native text; image content was not read.")
        page_count = document.page_count
    return pages, tables, warnings, page_count


def _docx(data):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    _check_office_archive(data)
    document = Document(io.BytesIO(data))
    pages, tables, warnings = [], [], []
    heading, previous, paragraph_number = "Document", "", 0

    def visit(container):
        nonlocal heading, previous, paragraph_number
        for item in container.iter_inner_content():
            if isinstance(item, Paragraph):
                text = item.text.strip()
                if not text:
                    continue
                paragraph_number += 1
                if item.style and item.style.name.startswith("Heading"):
                    heading = text
                pages.append(PageText(0, text, f"{heading}, paragraph {paragraph_number}"))
                previous = text
            elif isinstance(item, Table):
                number = len(tables) + 1
                parsed = make_table(
                    f"t{number}",
                    [[cell.text for cell in row.cells] for row in item.rows],
                    f"{heading}, Table {number}",
                    caption=previous,
                )
                if parsed:
                    tables.append(parsed)
                seen = set()
                for row in item.rows:
                    for cell in row.cells:
                        if cell._tc not in seen:
                            seen.add(cell._tc)
                            for nested in cell.tables:
                                number = len(tables) + 1
                                parsed = make_table(
                                    f"t{number}",
                                    [[c.text for c in r.cells] for r in nested.rows],
                                    f"{heading}, nested Table {number}",
                                )
                                if parsed:
                                    tables.append(parsed)

    visit(document)
    if document.inline_shapes:
        warnings.append("Embedded images/charts were not read; native text and tables were read.")
    if re.search(rb"<w:(ins|del)\b", document.part.blob):
        warnings.append("Tracked changes exist; accept/reject revisions in Word before indexing.")
    return pages, tables, warnings, 0


def _xlsx(data):
    from openpyxl import load_workbook

    _check_office_archive(data)
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    formulas = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    pages, tables, warnings = [], [], []
    try:
        if len(workbook.worksheets) > 50:
            raise ValidationError("Excel workbooks are limited to 50 worksheets.")
        for sheet, formula_sheet in zip(workbook.worksheets, formulas.worksheets, strict=True):
            # Valid streaming-generated XLSX files may omit <dimension>; their
            # read-only max_row/max_column are None. Bound the actual stream too.
            if (sheet.max_row or 0) * max(sheet.max_column or 0, 1) > MAX_WORKSHEET_CELLS:
                raise ValidationError(
                    f"Worksheet {sheet.title} exceeds the {MAX_WORKSHEET_CELLS:,}-cell "
                    "processing limit."
                )
            cached_rows = sheet.iter_rows(values_only=False)
            formula_rows = formula_sheet.iter_rows(values_only=False)
            matrix, row_numbers, formula_without_cache, width = [], [], 0, 0
            for row_number, (cached_row, formula_row) in enumerate(
                zip(cached_rows, formula_rows, strict=True), 1
            ):
                width = max(width, len(cached_row), len(formula_row))
                if row_number * width > MAX_WORKSHEET_CELLS:
                    raise ValidationError(
                        f"Worksheet {sheet.title} exceeds the {MAX_WORKSHEET_CELLS:,}-cell "
                        "processing limit."
                    )
                cells = []
                for cached, formula in zip(cached_row, formula_row, strict=True):
                    value = cached.value
                    if formula.data_type == "f" and value is None:
                        value = formula.value
                        formula_without_cache += 1
                    cells.append(_cell_text(value))
                while cells and not cells[-1]:
                    cells.pop()
                if any(cells):
                    matrix.append(cells)
                    row_numbers.append(row_number)
            if not matrix:
                continue
            width = max(map(len, matrix))
            matrix = [row + [""] * (width - len(row)) for row in matrix]
            location = f"Worksheet {sheet.title}"
            pages.append(PageText(0, f"Excel worksheet: {sheet.title}", location))
            # Keep report titles/notes as prose; the first populated row with
            # multiple fields is the table header, not a merged title in A1.
            header_index = next(
                (i for i, row in enumerate(matrix) if sum(bool(cell) for cell in row) >= 2),
                None,
            )
            if header_index is None and width == 1 and len(matrix) > 1:
                header_index = 0
            if header_index is not None and header_index < len(matrix) - 1:
                for row_number, row in zip(
                    row_numbers[:header_index], matrix[:header_index], strict=True
                ):
                    pages.append(
                        PageText(
                            0,
                            " | ".join(cell for cell in row if cell),
                            f"{location}, row {row_number}",
                        )
                    )
                caption = "\n".join(
                    " | ".join(cell for cell in row if cell) for row in matrix[:header_index]
                )
                table = make_table(
                    f"{sheet.title}-t1",
                    matrix[header_index:],
                    f"{location}, rows {row_numbers[header_index]}-{row_numbers[-1]}",
                    caption=f"Excel workbook worksheet: {sheet.title}\n{caption}".strip(),
                )
                if table:
                    tables.append(table)
            else:
                for row_number, row in zip(row_numbers, matrix, strict=True):
                    pages.append(
                        PageText(
                            0,
                            " | ".join(cell for cell in row if cell),
                            f"{location}, row {row_number}",
                        )
                    )
            if formula_without_cache:
                warnings.append(
                    f"{sheet.title}: {formula_without_cache} formula cells have no saved result; "
                    "their formula text was indexed. Recalculate and save in Excel for values."
                )
            if getattr(sheet, "_charts", ()):
                warnings.append(
                    f"{sheet.title}: embedded Excel chart appearance is not read; source cell "
                    "data is indexed when present."
                )
    finally:
        workbook.close()
        formulas.close()
    return pages, tables, warnings, 0


def _pptx(data):
    from pptx import Presentation

    _check_office_archive(data)
    presentation = Presentation(io.BytesIO(data))
    pages, tables, warnings = [], [], []
    for slide_number, slide in enumerate(presentation.slides, 1):
        prose = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text.strip()
                if text:
                    prose.append(text)
            if shape.has_table:
                ppt_table = shape.table
                parsed = make_table(
                    f"s{slide_number}-t{len(tables) + 1}",
                    [[cell.text for cell in row.cells] for row in ppt_table.rows],
                    f"Slide {slide_number}, Table {len(tables) + 1}",
                    slide_number,
                )
                if parsed:
                    tables.append(parsed)
            if shape.has_chart:
                chart = shape.chart
                for plot_number, plot in enumerate(chart.plots, 1):
                    try:
                        category_collection = plot.categories
                        try:
                            categories = [
                                " / ".join(str(part) for part in category)
                                for category in category_collection.flattened_labels
                            ]
                        except (AttributeError, TypeError):
                            categories = [str(category.label) for category in category_collection]
                        chart_series = list(plot.series)
                        headers = [
                            "Category",
                            *[
                                str(item.name or f"Series {i}")
                                for i, item in enumerate(chart_series, 1)
                            ],
                        ]
                        rows = []
                        for i, category in enumerate(categories):
                            values = [category]
                            for item in chart_series:
                                observations = list(item.values)
                                values.append(
                                    _cell_text(observations[i]) if i < len(observations) else ""
                                )
                            rows.append(values)
                    except (AttributeError, TypeError, ValueError) as exc:
                        warnings.append(
                            f"Slide {slide_number}: chart {plot_number} could not be read "
                            f"({type(exc).__name__}); nearby slide text was retained."
                        )
                        continue
                    parsed = make_table(
                        f"s{slide_number}-chart{len(tables) + 1}",
                        rows,
                        f"Slide {slide_number}, Chart {plot_number}",
                        slide_number,
                        caption=chart.chart_title.text_frame.text if chart.has_title else "",
                        headers=headers,
                    )
                    if parsed:
                        tables.append(parsed)
        if prose:
            pages.append(PageText(slide_number, "\n".join(prose), f"Slide {slide_number}"))
        if getattr(slide, "has_notes_slide", False):
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                pages.append(PageText(slide_number, notes, f"Slide {slide_number}, speaker notes"))
        if not prose and not any(t.page_number == slide_number for t in tables):
            warnings.append(f"Slide {slide_number}: no native text or table content was found.")
    return pages, tables, warnings, len(presentation.slides)
