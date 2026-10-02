from __future__ import annotations

import io
from dataclasses import replace

import pytest
from docx import Document

from rag_pdf.benchmark import configurations
from rag_pdf.calculations import calculate_table
from rag_pdf.chunking import chunk_document
from rag_pdf.config import ChunkConfig
from rag_pdf.document_processing import extract_document
from rag_pdf.embeddings import LocalEmbeddingModel
from rag_pdf.errors import ValidationError
from rag_pdf.models import ChatTurn, DocumentTable, ExtractedDocument, PageText, SearchResult
from rag_pdf.service import (
    ChatRepository,
    answer_question,
    expand_tables,
    index_document,
    library_id,
    search_library,
)
from rag_pdf.vector_store import create_store


class WordTokenizer:
    def encode(self, text, add_special_tokens=False):
        return text.split()


class FakeEmbedder:
    model_name = "test-model"
    tokenizer = WordTokenizer()
    batch_size = 2

    def embed_documents(self, texts):
        return [[1.0, 0.0] if "internship" in t else [0.0, 1.0] for t in texts]

    def embed_query(self, text):
        self.last_query = text
        return [1.0, 0.0] if "internship" in text else [0.0, 1.0]


def document(name="program.pdf", text="The internship lasts eight weeks."):
    return ExtractedDocument(name, name, 1, (PageText(1, text, "p. 1"),), "en", 1, len(text))


@pytest.mark.parametrize("backend", ["Chroma", "Qdrant", "LanceDB"])
def test_backend_library_search_and_persistence(tmp_path, backend):
    store = create_store(backend, tmp_path)
    embedder = FakeEmbedder()
    a, b = document(), document("support.docx", "The support desk is open on Mondays.")
    first = index_document(a, ChunkConfig(), embedder, store)
    second = index_document(b, ChunkConfig(), embedder, store)
    assert index_document(a, ChunkConfig(), embedder, store).reused_existing_index
    results = search_library([first.index_id, second.index_id], [1.0, 0.0], 2, store)
    assert results[0].filename == "program.pdf"
    assert results[0].location == "p. 1"
    assert results[0].score == pytest.approx(1.0, abs=1e-5)
    assert {r.filename for r in results} == {a.filename, b.filename}
    # Removing a document from the active library must exclude its retained index.
    assert all(
        r.filename == b.filename for r in search_library([second.index_id], [1.0, 0.0], 5, store)
    )
    if hasattr(store, "close"):
        store.close()
    reopened = create_store(backend, tmp_path)
    assert reopened.count(first.index_id) == first.chunk_count
    if hasattr(reopened, "close"):
        reopened.close()


def test_docx_paragraphs_tables_and_locations():
    source = Document()
    source.add_heading("Participation", level=1)
    source.add_paragraph("Revenue is recorded in USD millions.")
    table = source.add_table(rows=3, cols=2)
    for row, values in zip(
        table.rows,
        [["Year", "Revenue USD millions"], ["2024", "120"], ["2025", "150"]],
        strict=True,
    ):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    source.add_paragraph("The final audited values are shown above.")
    stream = io.BytesIO()
    source.save(stream)
    parsed = extract_document(stream.getvalue(), "report.docx")
    assert parsed.page_count == 0
    assert parsed.tables[0].headers == ("Year", "Revenue USD millions")
    assert parsed.tables[0].rows[-1] == ("2025", "150")
    assert "Participation" in parsed.tables[0].location
    assert parsed.tables[0].caption == "Revenue is recorded in USD millions."
    assert "audited" in parsed.pages[-1].text


def test_tables_split_with_headers_bounded_rows_and_overlap():
    table = DocumentTable(
        "t1", ("Month", "Visits"), tuple((f"Month{i}", str(i * 10)) for i in range(30)), "Table 1"
    )
    doc = replace(document(), pages=(), tables=(table,))
    chunks = chunk_document(doc, WordTokenizer(), ChunkConfig(32, 30))
    assert len(chunks) > 1
    assert all(c.text.startswith("Month | Visits") for c in chunks)
    assert all(c.token_count <= 32 for c in chunks)
    assert any(c.overlap_tokens > 0 for c in chunks)
    assert set(range(1, 31)) == {r for c in chunks for r in range(c.row_start, c.row_end + 1)}


def test_wide_table_row_is_not_silently_truncated():
    table = DocumentTable("t1", ("Description",), (("word " * 100,),), "Table 1")
    with pytest.raises(ValidationError, match="Increase the chunk size"):
        chunk_document(replace(document(), tables=(table,)), WordTokenizer(), ChunkConfig(32, 15))


def test_table_expansion_and_exact_calculation():
    table = DocumentTable("t1", ("Year", "Value"), (("2024", "120"), ("2025", "150")), "Table 1")
    doc = replace(document(), tables=(table,))
    hit = SearchResult("a", "partial", 1, 0, 1, doc.document_id, doc.filename, "", "t1")
    sources = expand_tables([hit, replace(hit, chunk_id="b")], [doc])
    assert len(sources) == 1
    assert "150" in sources[0].text
    result = calculate_table(
        sources, 1, "percent_change", [{"row": 1, "column": 2}, {"row": 2, "column": 2}]
    )
    assert result["result"] == "25.00"
    with pytest.raises(ValueError, match="coordinates"):
        calculate_table(sources, 1, "sum", [{"row": 0, "column": 2}])
    partial = expand_tables([hit], [doc], max_characters=100)
    # A cap that is too small cannot accidentally expose a complete-table calculator.
    tiny = expand_tables([hit], [doc], max_characters=5)
    assert not tiny or tiny[0].table is None
    assert partial


def test_follow_up_rewritten_before_retrieval_and_cited(tmp_path):
    store = create_store("Chroma", tmp_path)
    embedder = FakeEmbedder()
    doc = document()
    summary = index_document(doc, ChunkConfig(), embedder, store)

    class LLM:
        def rewrite_question(self, question, history):
            assert question == "How long is it?"
            assert history[0].content == "Tell me about the internship."
            return "How long is the internship?"

        def answer(self, question, sources):
            assert sources[0].filename == "program.pdf"
            return "Eight weeks. [Source 1]"

    answer = answer_question(
        index_ids=[summary.index_id],
        documents=[doc],
        question="How long is it?",
        history=[ChatTurn("user", "Tell me about the internship.")],
        top_k=5,
        embedder=embedder,
        store=store,
        language_model=LLM(),
    )
    assert embedder.last_query == "How long is the internship?"
    assert "Eight weeks" in answer.text


def test_history_persists_and_document_sets_are_isolated(tmp_path):
    a, b = document(), document("other.pdf", "Other text.")
    repository = ChatRepository(tmp_path)
    key = library_id([a, b])
    assert key == library_id([b, a])
    messages = [{"role": "user", "content": "How long is it?"}]
    repository.save(key, messages)
    assert ChatRepository(tmp_path).load(key) == messages
    assert repository.load(library_id([a])) == []


def test_embedding_instructions_are_model_specific():
    model = object.__new__(LocalEmbeddingModel)
    model.model_name = "BAAI/bge-m3"
    assert model.format_text("question", query=True) == "question"
    model.model_name = "intfloat/e5-base-v2"
    assert model.format_text("text", query=False) == "passage: text"
    model.model_name = "Qwen/Qwen3-Embedding-0.6B"
    assert model.format_text("question", query=True).endswith("Query: question")


def test_study_has_eleven_unique_one_factor_configurations():
    configs = configurations()
    assert len(configs) == len(set(configs)) == 11
    assert {c.overlap for c in configs} == {15, 30}
    for config in configs[1:]:
        fields = ["database", "embedding", "chunk_size", "overlap", "llm", "top_k"]
        assert sum(getattr(config, f) != getattr(configs[0], f) for f in fields) == 1
