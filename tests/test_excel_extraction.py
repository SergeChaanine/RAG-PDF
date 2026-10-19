"""Regression cases for report-style spreadsheets with missing dimensions."""

import io
import re
import zipfile

import pytest
from openpyxl import Workbook
from test_pipeline import FakeEmbedder, document
from test_retrieval import CrowdedStore

from rag_pdf.calculations import calculate_table
from rag_pdf.config import ChunkConfig
from rag_pdf.document_processing import extract_document
from rag_pdf.errors import ValidationError
from rag_pdf.service import answer_question, index_document


def save_without_dimensions(workbook, *, cached_total=False):
    source, output = io.BytesIO(), io.BytesIO()
    workbook.save(source)
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(output, "w") as rewritten:
        for name in original.namelist():
            content = original.read(name)
            if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                content = re.sub(rb"<dimension\b[^>]*/>", b"", content)
                if cached_total:
                    content = content.replace(
                        b"<f>SUM(B5:B16)</f><v></v>", b"<f>SUM(B5:B16)</f><v>220</v>"
                    )
            rewritten.writestr(name, content)
    return output.getvalue()


def bookshop_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Monthly Summary"
    sheet.append(["NORTHSTAR BOOKSHOP | MONTHLY PERFORMANCE"])
    sheet.append(["Reporting year: 2025. Fictional sales data."])
    sheet.append([])
    sheet.append(["Month", "Orders"])
    for month, count in enumerate([13, 14, 15, 14, 16, 17, 16, 17, 20, 22, 26, 30], 1):
        sheet.append([f"2025-{month:02d}", count])
    sheet.append(["2025 TOTAL", "=SUM(B5:B16)"])
    return workbook


def test_missing_dimensions_and_preamble_preserve_headers_values_and_cell_locations():
    data = save_without_dimensions(bookshop_workbook(), cached_total=True)
    doc = extract_document(data, "Northstar_Bookshop_Statistics_English.xlsx")
    table = doc.tables[0]
    assert table.headers == ("Month", "Orders")
    assert len(table.rows) == 13
    assert table.rows[-1] == ("2025 TOTAL", "220")
    assert table.location == "Worksheet Monthly Summary, rows 4-17"
    assert "NORTHSTAR BOOKSHOP" in table.caption
    assert any(page.location.endswith("row 1") for page in doc.pages)
    assert sum(int(row[1]) for row in table.rows[:12]) == 220
    assert not doc.warnings


def test_uncached_formulas_are_exposed_without_inventing_values():
    doc = extract_document(save_without_dimensions(bookshop_workbook()), "report.xlsx")
    assert doc.tables[0].rows[-1][1] == "=SUM(B5:B16)"
    assert any("no saved result" in warning for warning in doc.warnings)


def test_actual_stream_is_bounded_when_dimensions_are_missing(monkeypatch):
    monkeypatch.setattr("rag_pdf.document_processing.MAX_WORKSHEET_CELLS", 20)
    with pytest.raises(ValidationError, match="cell processing limit"):
        extract_document(save_without_dimensions(bookshop_workbook()), "large.xlsx")


def test_single_column_worksheets_keep_zero_and_negative_values():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Count"])
    sheet.append([0])
    sheet.append([-2])
    doc = extract_document(save_without_dimensions(workbook), "counts.xlsx")
    assert doc.tables[0].headers == ("Count",)
    assert doc.tables[0].rows == (("0",), ("-2",))


def test_exact_bookshop_question_retrieves_excel_and_uses_total_without_double_counting(tmp_path):
    excel = extract_document(
        save_without_dimensions(bookshop_workbook(), cached_total=True),
        "Northstar_Bookshop_Statistics_English.xlsx",
    )
    arabic = document("Arabic_RAG_Solar_Energy_Report.pdf", "معلومات الطاقة الشمسية 2025 " * 100)
    store, embedder = CrowdedStore(tmp_path), FakeEmbedder()
    summaries = [index_document(doc, ChunkConfig(), embedder, store) for doc in (excel, arabic)]

    class LLM:
        def prepare_search(self, question, history, filenames):
            return question, [question]

        def retrieval_queries(self, question, filenames):
            return [question]

        def rewrite_question(self, question, history):
            return question

        def answer(self, question, sources, *, original_question=None, visualization="auto"):
            source_number, source = next(
                (i, source) for i, source in enumerate(sources, 1) if source.table
            )
            assert source.filename == excel.filename
            assert source.table.headers == ("Month", "Orders")
            detail = calculate_table(
                sources, source_number, "sum",
                [{"row": i, "column": 2} for i in range(1, 13)],
            )
            total = calculate_table(sources, source_number, "sum", [{"row": 13, "column": 2}])
            assert detail["result"] == total["result"] == "220"
            return (
                f"| Year | Orders | Source |\n| --- | --- | --- |\n"
                f"| 2025 | 220 | [Source {source_number}] |"
            )

    answer = answer_question(
        index_ids=[summary.index_id for summary in summaries], documents=[excel, arabic],
        question="In northstar Bookshop, what is the total number of orders in 2025?",
        history=[], top_k=5, embedder=embedder, store=store, language_model=LLM(),
    )
    assert "220" in answer.text
    assert {source.filename for source in answer.sources} == {excel.filename, arabic.filename}
