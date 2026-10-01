from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

from rag_pdf.errors import ValidationError
from rag_pdf.evaluation import score_answer, token_f1
from rag_pdf.llm import GroqLanguageModel
from rag_pdf.memory_store import InMemoryVectorStore
from rag_pdf.models import (
    ChatTurn,
    ExtractedDocument,
    PageText,
    SearchResult,
    TextChunk,
)
from rag_pdf.pdf_processing import extract_document, extract_docx
from rag_pdf.service import answer_question
from rag_pdf.sqlite_store import SQLiteVectorStore


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph(
        "This quarterly financial report summarizes the company's performance. "
        "Revenue and profit are reported below as numeric statistics for the period. "
        "The figures are measured in thousands of dollars and compare business results."
    )
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Value (thousands USD)"
    row = table.add_row().cells
    row[0].text = "Revenue"
    row[1].text = "1250"
    row = table.add_row().cells
    row[0].text = "Profit"
    row[1].text = "300"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def make_chunk(
    chunk_id: str,
    filename: str,
    text: str,
    *,
    document_id: str = "doc-id",
) -> TextChunk:
    return TextChunk(
        chunk_id=chunk_id,
        document_id=document_id,
        filename=filename,
        text=text,
        page_number=1,
        chunk_index=0,
        page_chunk_index=0,
        token_count=2,
        start_index=0,
    )


def test_extract_docx_preserves_table_headers_with_values() -> None:
    document = extract_docx(make_docx(), "quarterly-results.docx")

    text = document.pages[0].text
    assert "Metric: Revenue" in text
    assert "Value (thousands USD): 1250" in text
    assert "Metric: Profit" in text
    assert "Value (thousands USD): 300" in text
    assert document.filename == "quarterly-results.docx"


def test_extract_document_routes_pdf_to_existing_extractor(monkeypatch: pytest.MonkeyPatch) -> None:
    expected = object()
    monkeypatch.setattr(
        "rag_pdf.pdf_processing.extract_pdf",
        lambda data, filename, rules: expected,
    )

    result = extract_document(b"pdf bytes", "report.pdf")

    assert result is expected


def test_extract_document_rejects_unsupported_file_types() -> None:
    with pytest.raises(ValidationError, match="Only PDF and Word"):
        extract_document(b"bytes", "table.xlsx")


def test_in_memory_store_ranks_vectors_and_keeps_source_filename() -> None:
    store = InMemoryVectorStore()
    document = ExtractedDocument(
        document_id="doc-id",
        filename="results.docx",
        page_count=1,
        pages=(PageText(page_number=1, text="Results"),),
        language="en",
        language_confidence=1.0,
        character_count=7,
    )
    chunks = [
        make_chunk("close", "results.docx", "close match"),
        make_chunk("far", "results.docx", "less similar"),
    ]
    store.add(
        "index",
        document,
        chunks,
        [[1.0, 0.0], [0.0, 1.0]],
        embedding_model="test-model",
        chunk_size=32,
        overlap_percent=15,
    )

    results = store.search("index", [1.0, 0.0], top_k=1)

    assert store.has_index("index")
    assert store.count("index") == 2
    assert results[0].chunk_id == "close"
    assert results[0].filename == "results.docx"
    assert results[0].score == pytest.approx(1.0)


def test_answer_question_merges_and_ranks_multiple_active_indexes() -> None:
    store = InMemoryVectorStore()
    document = ExtractedDocument(
        document_id="doc-id",
        filename="results.pdf",
        page_count=1,
        pages=(PageText(page_number=1, text="Results"),),
        language="en",
        language_confidence=1.0,
        character_count=7,
    )
    store.add(
        "index-a",
        document,
        [make_chunk("lower", "lower.pdf", "lower score")],
        [[0.6, 0.8]],
        embedding_model="test-model",
        chunk_size=32,
        overlap_percent=15,
    )
    store.add(
        "index-b",
        document,
        [make_chunk("higher", "higher.pdf", "higher score")],
        [[1.0, 0.0]],
        embedding_model="test-model",
        chunk_size=32,
        overlap_percent=15,
    )

    class Embedder:
        def embed_query(self, query: str) -> list[float]:
            return [1.0, 0.0]

    class LanguageModel:
        retrieved_sources: tuple[SearchResult, ...] = ()

        def rewrite_question(self, question: str, history: tuple[ChatTurn, ...]) -> str:
            return question

        def answer(self, question: str, sources: tuple[SearchResult, ...]) -> str:
            self.retrieved_sources = tuple(sources)
            return "Grounded answer"

    language_model = LanguageModel()
    response = answer_question(
        index_id=["index-a", "index-b"],
        question="Which result is better?",
        history=(),
        top_k=2,
        embedder=Embedder(),  # type: ignore[arg-type]
        store=store,
        language_model=language_model,  # type: ignore[arg-type]
    )

    assert response.text == "Grounded answer"
    assert [source.filename for source in language_model.retrieved_sources] == [
        "higher.pdf",
        "lower.pdf",
    ]


def test_answer_prompt_labels_document_sources_and_pdf_pages(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    model = object.__new__(GroqLanguageModel)
    captured: dict[str, str] = {}

    def capture_completion(**kwargs: str) -> str:
        captured.update(kwargs)
        return "Grounded answer"

    monkeypatch.setattr(model, "_complete", capture_completion)
    sources = [
        SearchResult("pdf-chunk", "PDF evidence", "annual.pdf", 4, 0, 0.9),
        SearchResult("word-chunk", "Word evidence", "results.docx", 1, 0, 0.8),
    ]

    result = model.answer("What happened?", sources)

    assert result == "Grounded answer"
    assert "annual.pdf | page 4" in captured["user"]
    assert "results.docx]" in captured["user"]
    assert "[filename, p. X]" in captured["user"]


def test_sqlite_store_persists_embeddings_and_ranks_by_cosine(tmp_path) -> None:
    database_path = tmp_path / "vectors.sqlite3"
    store = SQLiteVectorStore(database_path)
    document = ExtractedDocument(
        document_id="doc-id",
        filename="results.pdf",
        page_count=1,
        pages=(PageText(page_number=1, text="Results"),),
        language="en",
        language_confidence=1.0,
        character_count=7,
    )
    store.add(
        "index",
        document,
        [
            make_chunk("close", "results.pdf", "close match"),
            make_chunk("far", "results.pdf", "far match"),
        ],
        [[1.0, 0.0], [0.0, 1.0]],
        embedding_model="test-model",
        chunk_size=32,
        overlap_percent=15,
    )
    store.close()

    reopened = SQLiteVectorStore(database_path)
    results = reopened.search("index", [1.0, 0.0], top_k=2)

    assert reopened.count("index") == 2
    assert [result.chunk_id for result in results] == ["close", "far"]
    assert results[0].filename == "results.pdf"
    assert results[0].score == pytest.approx(1.0)
    reopened.close()


def test_token_f1_ignores_citations_but_measures_reference_overlap() -> None:
    assert token_f1("Revenue was 1250", "Revenue is 1250 [results.pdf, p. 2]") == pytest.approx(
        2 / 3
    )
    assert token_f1("same answer", "Same answer") == pytest.approx(1.0)


def test_unknown_case_counts_explicit_abstention_as_correct() -> None:
    result = score_answer(
        "I could not find this information in the selected documents.",
        should_abstain=True,
    )

    assert result["abstained"] is True
    assert result["outcome_correct"] is True
    assert result["outcome_label"] == "Correct abstention"


def test_unknown_case_rejects_an_unsupported_answer() -> None:
    result = score_answer("The answer is 42.", should_abstain=True)

    assert result["abstained"] is False
    assert result["outcome_correct"] is False
    assert result["outcome_label"] == "Unsupported answer"


def test_known_case_requires_answer_and_f1_threshold() -> None:
    good = score_answer("Revenue was 1250.", expected_answer="Revenue was 1250")
    refusal = score_answer("I don't know.", expected_answer="Revenue was 1250")

    assert good["outcome_correct"] is True
    assert good["token_f1"] == pytest.approx(1.0)
    assert refusal["outcome_correct"] is False
    assert refusal["outcome_label"] == "Incorrect abstention"
