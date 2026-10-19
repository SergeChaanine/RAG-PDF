"""Regression cases for mixed-language libraries and crowded vector rankings."""

import io
from dataclasses import replace

import pytest
from openpyxl import Workbook
from test_llm import model_with_client, response
from test_pipeline import FakeEmbedder, document

from rag_pdf.config import ChunkConfig
from rag_pdf.document_processing import extract_document
from rag_pdf.lexical import KeywordIndex
from rag_pdf.models import DocumentTable, PageText
from rag_pdf.service import answer_question, expand_tables, index_document, search_library
from rag_pdf.vector_store import result


class CrowdedStore:
    """Simulate a PDF scoring above every spreadsheet chunk, even on Excel queries."""

    def __init__(self, path):
        self.path = path
        self.chunks = {}

    def count(self, index_id):
        return len(self.chunks.get(index_id, []))

    def delete(self, index_id):
        self.chunks.pop(index_id, None)

    def add(self, index_id, document, chunks, embeddings):
        from dataclasses import asdict

        self.chunks.setdefault(index_id, []).extend(asdict(chunk) for chunk in chunks)

    def search(self, index_id, query_embedding, top_k):
        return [
            result(row, 0.99 if row["filename"].endswith(".pdf") else 0.01)
            for row in self.chunks[index_id][:top_k]
        ]


def mixed_library(tmp_path):
    arabic = document("policies.pdf", "سياسة الحضور والاجازات للطلاب. " * 100)
    workbook = Workbook()
    # Put the relevant sheet beyond the vector candidate limit, so broadening
    # semantic search alone cannot pass this case.
    for i in range(25):
        sheet = workbook.active if i == 0 else workbook.create_sheet()
        sheet.title = f"Inventory {i}"
        sheet.append(["Item", "Stock"])
        sheet.append(["Pen", i + 1])
    sheet = workbook.create_sheet("Payroll")
    sheet.append(["Employee", "Salary USD"])
    sheet.append(["Alice", 1200])
    sheet.append(["Bob", 1800])
    sheet = workbook.create_sheet("Monthly Summary")
    sheet.append(["Year", "Gross Profit"])
    sheet.append(["2025 TOTAL", 4351.539])
    data = io.BytesIO()
    workbook.save(data)
    excel = extract_document(data.getvalue(), "accounts.xlsx")
    store = CrowdedStore(tmp_path)
    embedder = FakeEmbedder()
    summaries = [index_document(doc, ChunkConfig(), embedder, store) for doc in (arabic, excel)]
    return (arabic, excel), store, embedder, summaries


@pytest.mark.parametrize(
    "question",
    [
        "What are the employee salaries on the Payroll worksheet?",
        "What is in the Payroll worksheet?",
        "What salaries are in the Excel file?",
    ],
)
def test_excel_answer_finds_sheet_even_when_arabic_pdf_dominates_vectors(tmp_path, question):
    documents, store, embedder, summaries = mixed_library(tmp_path)

    class LLM:
        def prepare_search(self, question, history, filenames):
            return question, [question]

        def retrieval_queries(self, question, filenames):
            return [question]

        def rewrite_question(self, question, history):
            return question

        def answer(self, question, sources, *, original_question=None, visualization="auto"):
            assert {source.filename for source in sources} == {"policies.pdf", "accounts.xlsx"}
            payroll = sources[0]
            assert payroll.filename == "accounts.xlsx"
            assert "Payroll" in payroll.location
            assert payroll.table.headers == ("Employee", "Salary USD")
            assert payroll.table.rows == (("Alice", "1200"), ("Bob", "1800"))
            assert payroll.retrieval_method == "keyword"
            return (
                "| Employee | Salary USD | Source |\n| --- | --- | --- |\n"
                "| Alice | 1200 | [Source 1] |\n| Bob | 1800 | [Source 1] |"
            )

    answer = answer_question(
        index_ids=[summary.index_id for summary in summaries],
        documents=documents,
        question=question,
        history=[],
        top_k=2,
        embedder=embedder,
        store=store,
        language_model=LLM(),
    )
    assert "1800" in answer.text


def test_each_document_gets_a_slot_even_without_keyword_overlap(tmp_path):
    _, store, _, summaries = mixed_library(tmp_path)
    hits = search_library([summary.index_id for summary in summaries], [0, 1], 2, store)
    assert len(hits) == 2
    assert {hit.filename for hit in hits} == {"policies.pdf", "accounts.xlsx"}
    # Coverage still respects the requested limit if there are too few slots.
    assert len(search_library([summary.index_id for summary in summaries], [0, 1], 1, store)) == 1


def test_bilingual_candidates_keep_two_distinct_results_per_document_when_room(tmp_path):
    _, store, _, summaries = mixed_library(tmp_path)
    hits = search_library(
        [s.index_id for s in summaries], [0, 1], 4, store,
        question="ما إجمالي الربح؟", query_variants=[("What is the total profit?", [0, 1])],
    )
    assert len(hits) == 4
    assert sum(hit.filename == "accounts.xlsx" for hit in hits) == 2
    assert sum(hit.filename == "policies.pdf" for hit in hits) == 2
    assert len({(hit.document_id, hit.chunk_id) for hit in hits}) == 4


def test_exact_arabic_profit_question_finds_english_sheet_beyond_vector_candidates(tmp_path):
    documents, store, embedder, summaries = mixed_library(tmp_path)
    question = "ما إجمالي الربح الذي حققته المكتبة خلال عام 2025؟"
    translated = "What is the total profit earned by the bookshop in 2025?"
    expected = (
        "| السنة | إجمالي الربح | المصدر |\n| --- | --- | --- |\n"
        "| 2025 | 4351.539 | [Source 1] |"
    )
    model = model_with_client()
    model._client.chat.completions.create.side_effect = [response(translated), response(expected)]

    answer = answer_question(
        index_ids=[s.index_id for s in summaries], documents=documents, question=question,
        history=[], top_k=2, embedder=embedder, store=store, language_model=model,
    )

    assert answer.text == expected
    assert len(answer.sources) == 2
    assert answer.sources[0].filename == "accounts.xlsx"
    assert "Monthly Summary" in answer.sources[0].location
    assert answer.sources[0].table.rows == (("2025 TOTAL", "4351.539"),)
    assert answer.standalone_question == question
    answer_prompt = model._client.chat.completions.create.call_args.kwargs["messages"]
    assert "Required response language: Arabic" in answer_prompt[0]["content"]


def test_english_question_finds_arabic_evidence_beyond_vector_candidates(tmp_path):
    arabic = replace(
        document("solar.pdf"),
        pages=tuple(PageText(i, "معلومات عامة عن الطاقة", f"p. {i}") for i in range(1, 26))
        + (PageText(26, "اعتماد خطة التوسعة في 18 مارس 2025", "p. 26"),),
    )
    store, embedder = CrowdedStore(tmp_path), FakeEmbedder()
    summary = index_document(arabic, ChunkConfig(), embedder, store)
    question = "On what date was the expansion plan approved?"
    expected = "The expansion plan was approved on 18 March 2025. [Source 1]"
    model = model_with_client()
    model._client.chat.completions.create.side_effect = [
        response("في أي تاريخ تم اعتماد خطة التوسعة؟"), response(expected)
    ]

    answer = answer_question(
        index_ids=[summary.index_id], documents=[arabic], question=question, history=[],
        top_k=1, embedder=embedder, store=store, language_model=model,
    )

    assert answer.text == expected
    assert answer.sources[0].page_number == 26
    assert answer.sources[0].retrieval_method == "keyword"


def test_keyword_search_uses_only_active_documents_and_finds_named_sheet(tmp_path):
    documents, store, _, summaries = mixed_library(tmp_path)
    keyword_path = store.path / "keywords.sqlite3"
    # A reused vector index must also repair a missing keyword index.
    keyword_path.unlink()
    excel = documents[1]
    assert index_document(excel, ChunkConfig(), FakeEmbedder(), store).reused_existing_index
    keywords = KeywordIndex(keyword_path)
    hits = keywords.search([summaries[1].index_id], 'Show "Payroll" OR salary?', 5)
    assert hits and all(hit.filename == "accounts.xlsx" for hit in hits)
    assert any(hit.table_id == "Payroll-t1" for hit in hits)
    hits = search_library(
        [summaries[0].index_id], [0, 1], 5, store, question="Payroll employee salary"
    )
    assert hits and all(hit.filename == "policies.pdf" for hit in hits)


def test_arabic_keyword_matching_normalizes_diacritics_and_alef(tmp_path):
    doc = document("report.pdf", "إحصاءات الإيرادات للسنة الحالية")
    store = CrowdedStore(tmp_path)
    summary = index_document(doc, ChunkConfig(), FakeEmbedder(), store)
    hits = KeywordIndex(store.path / "keywords.sqlite3").search(
        [summary.index_id], "ما هي احصاءات الايرادات؟", 5
    )
    assert hits and hits[0].text == doc.pages[0].text


def test_table_chunks_do_not_fill_slots_needed_by_other_worksheets(tmp_path):
    doc = replace(
        document("numbers.xlsx"),
        pages=(),
        tables=(
            DocumentTable("a", ("Product", "Sales"), tuple((str(i), "2") for i in range(40)), "A"),
            DocumentTable("b", ("Product", "Sales"), (("Other", "10"),), "B"),
        ),
    )
    store = CrowdedStore(tmp_path)
    summary = index_document(doc, ChunkConfig(32, 15), FakeEmbedder(), store)
    hits = search_library([summary.index_id], [0, 1], 5, store, question="Product sales")
    assert {hit.table_id for hit in hits} == {"a", "b"}
    assert len(hits) == 2


def test_large_table_cannot_consume_another_documents_context_budget(tmp_path):
    documents, store, _, summaries = mixed_library(tmp_path)
    hits = search_library(
        [summary.index_id for summary in summaries], [0, 1], 2, store,
        question="Payroll employee salary",
    )
    pdf_hit = next(hit for hit in hits if hit.filename == "policies.pdf")
    excel_hit = next(hit for hit in hits if hit.filename == "accounts.xlsx")
    huge_table = DocumentTable("huge", ("Text",), (("x" * 39000,),), "p. 1")
    arabic = replace(documents[0], tables=(huge_table,))
    pdf_hit = replace(pdf_hit, table_id="huge", text="A short table excerpt.")
    sources = expand_tables([pdf_hit, excel_hit], [arabic, documents[1]])
    assert len(sources) == 2
    assert sources[0].table is None and "PARTIAL TABLE" in sources[0].text
    assert sources[1].table.rows[-1] == ("Bob", "1800")


def test_invalid_top_k_is_rejected(tmp_path):
    with pytest.raises(ValueError, match="Top K"):
        search_library([], [0, 1], 0, CrowdedStore(tmp_path))
